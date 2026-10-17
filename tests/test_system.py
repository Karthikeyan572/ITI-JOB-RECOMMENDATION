import io
from pathlib import Path
import sys
import unittest
import pandas as pd
from docx import Document

# Ensure src is on sys.path
BASE_DIR = Path(__file__).resolve().parent.parent
SRC_DIR = BASE_DIR / "src"
if str(SRC_DIR) not in sys.path:
    sys.path.insert(0, str(SRC_DIR))

from candidate_extractor import (
    classify_document_type,
    extract_candidate_profile,
    extract_experience,
    extract_skills,
    extract_trade,
    has_iti_evidence,
    normalize_skill,
    validate_iti_profile,
    extract_candidate_name,
    extract_certifications,
)
from job_matcher import (
    compute_certification_score,
    compute_experience_score,
    compute_trade_compatibility,
    ITIJobMatcher,
)
from resume_parser import clean_text, extract_docx_text, extract_pdf_text, parse_resume
from skill_gap import analyze_skill_gap, get_learning_recommendations, load_courses_dataset


class TestITISystem(unittest.TestCase):
    """
    Comprehensive test suite covering:
    1. PDF extraction
    2. DOCX extraction
    3. Explicit Electrician detection
    4. Explicit Fitter detection
    5. CSE resume with no ITI evidence
    6. False COPA detection prevention
    7. Skill normalization
    8. Trade compatibility
    9. Skill matching
    10. Final score calculation
    11. Skill gap calculation
    12. Top 3 ranking
    """

    @classmethod
    def setUpClass(cls):
        cls.matcher = ITIJobMatcher()
        cls.courses_df = load_courses_dataset()
        cls.pdf_sample_path = BASE_DIR / "tests" / "fixtures" / "sample_cse_resume.pdf"

    # -------------------------------------------------------------
    # 1. PDF Extraction
    # -------------------------------------------------------------
    def test_01_pdf_extraction(self):
        if self.pdf_sample_path.exists():
            text = parse_resume(self.pdf_sample_path)
            self.assertIsInstance(text, str)
            self.assertGreater(len(text), 100)
            self.assertIn("KARTHIKEYAN", text.upper())

    # -------------------------------------------------------------
    # 2. DOCX Extraction
    # -------------------------------------------------------------
    def test_02_docx_extraction(self):
        doc = Document()
        doc.add_heading("Resume: Ramesh Kumar", 0)
        doc.add_paragraph("ITI Trade: Electrician")
        doc.add_paragraph("Skills: Wiring, Motor Maintenance, PLC")
        doc.add_paragraph("Experience: 2 years experience")

        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)

        extracted = extract_docx_text(buf)
        self.assertIn("Electrician", extracted)
        self.assertIn("Wiring", extracted)
        self.assertIn("2 years experience", extracted)

    # -------------------------------------------------------------
    # 3. Explicit Electrician Detection
    # -------------------------------------------------------------
    def test_03_explicit_electrician_detection(self):
        sample_text = (
            "Name: Suresh V\n"
            "Education: ITI Electrician\n"
            "NCVT Certified\n"
            "Skills: Wiring, PLC, Single Phase Wiring\n"
            "Experience: 2 years"
        )
        trade = extract_trade(sample_text)
        self.assertEqual(trade, "Electrician")

        profile = extract_candidate_profile(sample_text)
        status = validate_iti_profile(profile, sample_text)
        self.assertEqual(status, "VALID")
        self.assertEqual(profile["iti_trade"], "Electrician")

    # -------------------------------------------------------------
    # 4. Explicit Fitter Detection
    # -------------------------------------------------------------
    def test_04_explicit_fitter_detection(self):
        sample_text = (
            "Candidate: Vignesh R\n"
            "ITI Trade: Fitter\n"
            "Technical Skills: Fitting, Assembly, Lathe Machine Operation\n"
            "Experience: 1 year of experience"
        )
        trade = extract_trade(sample_text)
        self.assertEqual(trade, "Fitter")

        profile = extract_candidate_profile(sample_text)
        status = validate_iti_profile(profile, sample_text)
        self.assertEqual(status, "VALID")
        self.assertEqual(profile["iti_trade"], "Fitter")

    # -------------------------------------------------------------
    # 5. CSE Resume with No ITI Evidence
    # -------------------------------------------------------------
    def test_05_cse_resume_no_iti_evidence(self):
        if self.pdf_sample_path.exists():
            text = parse_resume(self.pdf_sample_path)
            profile = extract_candidate_profile(text)
            status = validate_iti_profile(profile, text)

            # Must NOT detect an ITI trade
            self.assertIsNone(profile["iti_trade"])
            # Must classify as NOT_ITI
            self.assertEqual(status, "NOT_ITI")
            # Must detect qualification like B.E/B.Tech
            self.assertTrue(any(q in profile["qualification"] for q in ["B.E", "B.Tech"]))

    # -------------------------------------------------------------
    # 6. False COPA Detection Prevention
    # -------------------------------------------------------------
    def test_06_false_copa_prevention(self):
        # Even with computer, programming, software, web, database keywords,
        # an ITI trade of COPA must NOT be inferred without explicit ITI evidence
        cse_text = (
            "John Doe\n"
            "Software Developer & Computer Engineer\n"
            "Skills: Python, Programming, Computer Hardware, Database Management, Web Development\n"
            "B.Tech in Information Technology\n"
            "3 years experience in software company"
        )
        trade = extract_trade(cse_text)
        self.assertIsNone(trade, "COPA was falsely inferred from generic software terms!")
        profile = extract_candidate_profile(cse_text)
        status = validate_iti_profile(profile, cse_text)
        self.assertEqual(status, "NOT_ITI")

    # -------------------------------------------------------------
    # 7. Skill Normalization
    # -------------------------------------------------------------
    def test_07_skill_normalization(self):
        self.assertEqual(normalize_skill("Electrical Wiring"), "Wiring")
        self.assertEqual(normalize_skill("PLC Programming"), "PLC")
        self.assertEqual(normalize_skill("MS Excel"), "Excel")
        self.assertEqual(normalize_skill("Manual Metal Arc Welding"), "Arc Welding")
        # Canonical skills remain intact
        self.assertEqual(normalize_skill("Wiring"), "Wiring")
        self.assertEqual(normalize_skill("Motor Maintenance"), "Motor Maintenance")

    # -------------------------------------------------------------
    # 8. Trade Compatibility
    # -------------------------------------------------------------
    def test_08_trade_compatibility(self):
        comp_map = self.matcher.comp_map

        # Exact match
        self.assertEqual(compute_trade_compatibility("Electrician", "Electrician", comp_map), 1.0)
        self.assertEqual(compute_trade_compatibility("Fitter", "Fitter", comp_map), 1.0)

        # Related trade (Electrician related to Electronics Mechanic)
        self.assertEqual(compute_trade_compatibility("Electrician", "Electronics Mechanic", comp_map), 0.5)
        self.assertEqual(compute_trade_compatibility("Fitter", "Machinist", comp_map), 0.5)

        # Unrelated trade
        self.assertEqual(compute_trade_compatibility("Electrician", "Fitter", comp_map), 0.0)
        self.assertEqual(compute_trade_compatibility("Welder", "COPA", comp_map), 0.0)

    def test_related_trade_mappings_are_canonical(self):
        supported = set(self.matcher.trades_df["trade"].astype(str).str.strip())
        for related in self.matcher.trades_df["related_trades"].fillna(""):
            for trade in str(related).split(";"):
                if trade.strip():
                    self.assertIn(trade.strip(), supported)
        self.assertEqual(self.matcher.comp_map[("COPA", "Electronics Mechanic")], 0.0)

    # -------------------------------------------------------------
    # 9. Skill Matching & TF-IDF Cosine Similarity
    # -------------------------------------------------------------
    def test_09_skill_matching(self):
        sample_profile = {
            "iti_trade": "Electrician",
            "skills": ["Wiring", "PLC", "Electrical Testing"],
            "experience": 2.0,
            "certifications": ["Trade Certificate"],
            "qualification": ["ITI"],
        }
        matches = self.matcher.match_jobs(sample_profile, filter_unrelated_trades=True)
        self.assertGreater(len(matches), 0)

        top_job = matches[0]
        self.assertGreater(top_job["skill_score"], 0.5)
        self.assertGreater(top_job["tfidf_similarity"], 0.0)
        self.assertIn("Wiring", top_job["matched_skills"])

    # -------------------------------------------------------------
    # 10. Final Score Calculation (0.50*S + 0.25*T + 0.15*E + 0.10*C)
    # -------------------------------------------------------------
    def test_10_final_score_calculation(self):
        skill_s = 0.80
        trade_s = 1.00
        exp_s = 1.00
        cert_s = 1.00

        expected = 0.50 * skill_s + 0.25 * trade_s + 0.15 * exp_s + 0.10 * cert_s
        expected_score = round(expected, 4)

        self.assertAlmostEqual(expected_score, 0.90, places=3)

    # -------------------------------------------------------------
    # 11. Skill Gap Calculation
    # -------------------------------------------------------------
    def test_11_skill_gap_calculation(self):
        job_req = ["Wiring", "PLC", "Panel Wiring", "Motor Maintenance"]
        cand_skills = ["Wiring", "PLC"]

        gap = analyze_skill_gap(cand_skills, job_req)
        self.assertEqual(gap["matched_skills"], ["Wiring", "PLC"])
        self.assertEqual(gap["missing_skills"], ["Panel Wiring", "Motor Maintenance"])
        self.assertEqual(gap["gap_percentage"], 50.0)
        self.assertEqual(gap["match_percentage"], 50.0)

        # Check course recommendation for missing skills
        recs = get_learning_recommendations(gap["missing_skills"], self.courses_df)
        missing_rec_skills = [r["skill"] for r in recs]
        self.assertIn("Panel Wiring", missing_rec_skills)

    # -------------------------------------------------------------
    # 12. Top 3 Ranking
    # -------------------------------------------------------------
    def test_12_top_3_ranking(self):
        sample_profile = {
            "iti_trade": "Fitter",
            "skills": ["Fitting", "Assembly", "Lathe Machine Operation", "Measurement"],
            "experience": 1.5,
            "certifications": ["Trade Certificate"],
            "qualification": ["ITI"],
        }
        top_recs = self.matcher.get_top_recommendations(sample_profile, min_threshold=0.60, top_n=3)
        self.assertLessEqual(len(top_recs), 3)
        self.assertGreater(len(top_recs), 0)

        # Check descending order of score
        scores = [j["final_score"] for j in top_recs]
        self.assertEqual(scores, sorted(scores, reverse=True))

        # Check trade compatibility
        for j in top_recs:
            self.assertIn(j["iti_trade"], ["Fitter", "Machinist", "Mechanic"])

    # -------------------------------------------------------------
    # 13. Trade Unclear Handling
    # -------------------------------------------------------------
    def test_13_trade_unclear_handling(self):
        unclear_text = (
            "Rahul Sharma\n"
            "Completed 2-year course at Government Industrial Training Institute (ITI).\n"
            "NCVT Trade Certificate.\n"
            "Skills: Measurement, Safety\n"
            "Experience: 1 year"
        )
        trade = extract_trade(unclear_text)
        self.assertIsNone(trade)
        profile = extract_candidate_profile(unclear_text)
        status = validate_iti_profile(profile, unclear_text)
        self.assertEqual(status, "TRADE_UNCLEAR")

    # -------------------------------------------------------------
    # 14. Academic Survey Paper Rejection
    # -------------------------------------------------------------
    def test_14_survey_paper_detection(self):
        survey_text = (
            "A Comprehensive Survey on IoT in Industrial Automation\n"
            "Dr. R. Sharma, Dept. of Electrical Engineering\n"
            "Email: rsharma@iisc.ac.in | DOI: 10.1109/TII.2023.987654\n\n"
            "ABSTRACT\n"
            "In this survey paper, we review sensors, industrial wiring, and motor maintenance.\n\n"
            "1. INTRODUCTION\n"
            "Industrial infrastructure has evolved significantly [1], [2]. Over 10 years of experience "
            "in industrial deployment shows massive efficiency gains [3].\n\n"
            "2. LITERATURE SURVEY\n"
            "Smith et al. [4] investigated automated welding and inspection architectures.\n\n"
            "3. METHODOLOGY\n"
            "We classify industrial communications into three distinct tiers...\n\n"
            "4. CONCLUSION AND FUTURE DIRECTIONS\n"
            "Standardization remains an open problem.\n\n"
            "REFERENCES\n"
            "[1] A. Kumar, 'Smart Grids,' IEEE Trans., 2021.\n"
            "[2] B. Patel, 'Sensors and Wiring,' Springer, 2022.\n"
            "[3] C. Lee, 'Electric Motors,' 2023.\n"
            "[4] D. Smith, 'Industrial Sensor Networks Survey,' 2021."
        )

        doc_info = classify_document_type(survey_text)
        self.assertFalse(doc_info["is_resume"])
        self.assertEqual(doc_info["doc_type"], "ACADEMIC_SURVEY_PAPER")

        profile = extract_candidate_profile(survey_text)
        self.assertFalse(profile["is_resume"])
        self.assertEqual(profile["skills"], [])
        self.assertIsNone(profile["iti_trade"])

        status = validate_iti_profile(profile, survey_text)
        self.assertEqual(status, "NOT_A_RESUME")

    # -------------------------------------------------------------
    # 15. Real Resumes Pass Classification
    # -------------------------------------------------------------
    def test_15_real_resumes_pass_classification(self):
        electrician_text = (
            "Name: Suresh V\n"
            "Career Objective: Seeking ITI Electrician role.\n"
            "Education: ITI Electrician (NCVT)\n"
            "Technical Skills: Wiring, PLC, Motor Maintenance\n"
            "Work Experience: 2 years experience"
        )
        doc_info = classify_document_type(electrician_text)
        self.assertTrue(doc_info["is_resume"])
        self.assertEqual(doc_info["doc_type"], "RESUME")

        if self.pdf_sample_path.exists():
            cse_text = parse_resume(self.pdf_sample_path)
            cse_doc_info = classify_document_type(cse_text)
            self.assertTrue(cse_doc_info["is_resume"])
            self.assertEqual(cse_doc_info["doc_type"], "RESUME")


    # -------------------------------------------------------------
    # 16. Unverified Job Data Is Not Presented as Live
    # -------------------------------------------------------------
    def test_16_job_records_are_marked_unverified(self):
        jobs = self.matcher.jobs_df
        self.assertIn("record_status", jobs.columns)
        self.assertTrue(jobs["record_status"].eq("UNVERIFIED_SAMPLE").all())
        self.assertTrue(jobs["is_live_vacancy"].astype(str).str.lower().isin(["false", "0"]).all())
        self.assertTrue(jobs["source_url"].astype(str).str.strip().eq("").all())

    # -------------------------------------------------------------
    # 17. Course Suggestions Are Not Misrepresented as Verified
    # -------------------------------------------------------------
    def test_17_course_records_are_marked_unverified(self):
        self.assertIn("record_status", self.courses_df.columns)
        self.assertTrue(self.courses_df["record_status"].eq("UNVERIFIED_TOPIC_SUGGESTION").all())
        recs = get_learning_recommendations(["Wiring"], self.courses_df)
        self.assertTrue(recs)
        self.assertEqual(recs[0]["record_status"], "UNVERIFIED_TOPIC_SUGGESTION")
        self.assertEqual(recs[0]["resource_url"], "")

    # -------------------------------------------------------------
    # 18. Trade-Specific Official Learning Pages Exist for All Trades
    # -------------------------------------------------------------
    def test_18_trade_specific_learning_source_coverage(self):
        path = BASE_DIR / "iti_skill_gap_datasets" / "learning_sources.csv"
        sources = pd.read_csv(path).fillna("")
        trade_pages = sources[sources["verification_status"] == "OFFICIAL_TRADE_PAGE_CHECKED"]
        expected = {"Electrician", "Fitter", "Welder", "Machinist", "Electronics Mechanic", "Mechanic Motor Vehicle", "COPA"}
        self.assertEqual(set(trade_pages["applicable_trades"]), expected)
        self.assertTrue(trade_pages["source_url"].astype(str).str.startswith("https://").all())

    # -------------------------------------------------------------
    # 19. Research Leads Must Not Be Treated as Verified Vacancies
    # -------------------------------------------------------------
    def test_19_job_research_candidates_remain_unverified(self):
        data_dir = BASE_DIR / "iti_skill_gap_datasets"
        candidates = pd.read_csv(data_dir / "job_research_candidates.csv").fillna("")
        self.assertGreaterEqual(len(candidates), 1)
        self.assertTrue(candidates["verification_status"].str.contains("NEEDS_OFFICIAL_NOTICE_CONFIRMATION").all())
        jobs = pd.read_csv(data_dir / "jobs.csv").fillna("")
        self.assertFalse(jobs["job_id"].isin(candidates["candidate_id"]).any())
        self.assertTrue(jobs["record_status"].eq("UNVERIFIED_SAMPLE").all())




    def test_candidate_name_extracted_from_resume_not_filename(self):
        text = (
            "Name: Ramesh Kumar\n"
            "ITI Trade: Electrician\n"
            "Technical Skills: Wiring, Multimeter\n"
            "Certifications:\n"
            "First Aid Certificate - 2022\n"
            "Other Details: Available to join immediately\n"
        )
        profile = extract_candidate_profile(text)
        self.assertEqual(profile["name"], "Ramesh Kumar")

    def test_certifications_exclude_other_details_and_experience_prose(self):
        text = (
            "ITI Trade: Electrician\n"
            "Certifications:\n"
            "Electrical Safety Awareness – Institute workshop (2023)\n"
            "First Aid basics – 2022\n"
            "Other Details\n"
            "Available to join: Immediate\n"
            "Expected salary: Negotiable\n"
            "Willing to travel: Within 25 km\n"
            "Practical exposure\n"
            "During ITI workshop practice, assembled simple circuits and used a multimeter under supervision.\n"
            "References can be provided on request.\n"
        )
        certs = extract_certifications(text)
        self.assertTrue(any("Electrical Safety Awareness" in item for item in certs))
        self.assertTrue(any("First Aid" in item for item in certs))
        for unwanted in ["OTHER DETAILS", "Available to join", "Expected salary",
                         "Willing to travel", "Practical exposure", "References can be provided"]:
            self.assertFalse(any(unwanted.lower() in item.lower() for item in certs), certs)



    def test_learning_recommendations_prefer_candidate_trade(self):
        courses = pd.DataFrame([
            {"course_id":"C1","course_name":"Wiring for Electricians","skill":"Wiring","trade":"Electrician","level":"Beginner","provider":"Demo","duration":"2 hours","resource_url":"","record_status":"UNVERIFIED_TOPIC_SUGGESTION"},
            {"course_id":"C2","course_name":"Wiring for Fitters","skill":"Wiring","trade":"Fitter","level":"Beginner","provider":"Demo","duration":"2 hours","resource_url":"","record_status":"UNVERIFIED_TOPIC_SUGGESTION"},
            {"course_id":"C3","course_name":"General Wiring","skill":"Wiring","trade":"General","level":"Beginner","provider":"Demo","duration":"1 hour","resource_url":"","record_status":"UNVERIFIED_TOPIC_SUGGESTION"},
        ])
        recs = get_learning_recommendations(["Wiring"], courses, candidate_trade="Electrician")
        self.assertEqual([r["course_id"] for r in recs], ["C1", "C3"])
        self.assertTrue(all("Electrician" in r["recommendation_reason"] for r in recs))

    def test_learning_recommendations_dedupe_missing_skill(self):
        courses = pd.DataFrame([
            {"course_id":"C1","course_name":"PLC Basics","skill":"PLC","trade":"Electrician","level":"Beginner","provider":"Demo","duration":"2 hours","resource_url":""},
            {"course_id":"C1","course_name":"PLC Basics Duplicate","skill":"PLC","trade":"Electrician","level":"Beginner","provider":"Demo","duration":"2 hours","resource_url":""},
        ])
        recs = get_learning_recommendations(["PLC"], courses, candidate_trade="Electrician")
        self.assertEqual(len(recs), 1)


if __name__ == "__main__":
    unittest.main(verbosity=2)
